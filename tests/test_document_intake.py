"""文件导入：各类文件提取文字、文件头与压缩包检查、按密级控制 OCR、成员身份、不落盘。"""
import io
import os
import tempfile
import unittest
import uuid
import zipfile
from email.message import EmailMessage
from unittest.mock import AsyncMock, patch

from tests import _route_client as rc  # 先导入：测试用的异步连接不复用（必须在导入 models 之前）

from sqlalchemy import text

from models.init_db import SessionLocal
from service import document_intake as intake
from service.exceptions import InvalidInput
from tests._async_helpers import run_async
from tests.test_enterprise_access import _add_org_member, _add_team_member, _create_org, _create_team

_AVAILABLE, _WHY = rc.route_tests_available()


def make_docx() -> bytes:
    import docx
    doc = docx.Document()
    doc.add_paragraph("例会纪要：新版首页本月发布")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "负责人"
    table.rows[0].cells[1].text = "demo_emp"
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def make_xlsx() -> bytes:
    from openpyxl import Workbook
    book = Workbook()
    sheet = book.active
    sheet.title = "费用"
    sheet.append(["日期", "金额", "说明"])
    sheet.append(["2026-10-01", 260, "高铁票"])
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def make_pdf(text_: str = "Meeting minutes 2026 owner demo_emp", scanned: bool = False) -> bytes:
    import pymupdf
    doc = pymupdf.open()
    page = doc.new_page()
    if scanned:
        from PIL import Image
        picture = io.BytesIO()
        Image.new("RGB", (60, 60), "white").save(picture, "PNG")
        page.insert_image(pymupdf.Rect(72, 72, 200, 200), stream=picture.getvalue())
    else:
        page.insert_text((72, 72), text_)
    return doc.tobytes()


def make_png() -> bytes:
    from PIL import Image
    buffer = io.BytesIO()
    Image.new("RGB", (20, 20), "white").save(buffer, "PNG")
    return buffer.getvalue()


def make_eml(html_only: bool = False) -> bytes:
    message = EmailMessage()
    message["Subject"] = "关于下周发布的安排"
    message["From"] = "boss@example.com"
    message["To"] = "team@example.com"
    message["Date"] = "Mon, 05 Oct 2026 09:00:00 +0800"
    if html_only:
        message.set_content("<html><body><p>请 demo_emp 周五前提交方案</p><script>alert(1)</script><p>谢谢</p></body></html>", subtype="html")
    else:
        message.set_content("请 demo_emp 周五前提交方案。")
    message.add_attachment(b"fake", maintype="application", subtype="octet-stream", filename="方案.zip")
    return message.as_bytes()


class FileTypeTest(unittest.TestCase):
    def test_whitelist(self):
        self.assertEqual(intake.file_type_of("纪要.PDF"), "pdf")
        for bad in ("a.exe", "a.zip", "a.doc", "noext", "a.docx.exe"):
            with self.subTest(name=bad), self.assertRaises(InvalidInput):
                intake.file_type_of(bad)

    def test_signature_must_match_the_extension(self):
        intake._check_signature("pdf", make_pdf())
        intake._check_signature("docx", make_docx())
        intake._check_signature("png", make_png())
        intake._check_signature("txt", "你好".encode("utf-8"))
        for kind, content in (("pdf", b"just text"), ("docx", b"%PDF-1.4"), ("png", b"\xff\xd8\xff"), ("jpg", make_png()),
                              ("txt", b"MZ\x90\x00 binary"), ("txt", b"abc\x00def"), ("eml", b"\x7fELF....")):
            with self.subTest(kind=kind), self.assertRaises(InvalidInput):
                intake._check_signature(kind, content)

    def test_archive_bomb_is_rejected(self):
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("big.xml", b"0" * (90 * 1024 * 1024))
        self.assertLess(len(buffer.getvalue()), 1024 * 1024)
        with self.assertRaises(InvalidInput):
            intake._check_archive(buffer.getvalue())
        with self.assertRaises(InvalidInput):
            intake._check_archive(b"PK\x03\x04 not really a zip")


class ParseTest(unittest.TestCase):
    def parse(self, kind, content, ocr=False):
        return intake._parse_sync(kind, content, 1, ocr)

    def test_docx_keeps_paragraphs_and_tables(self):
        text_ = self.parse("docx", make_docx())["text"]
        self.assertIn("例会纪要：新版首页本月发布", text_)
        self.assertIn("demo_emp", text_)

    def test_xlsx_rows_become_pipe_separated_lines(self):
        text_ = self.parse("xlsx", make_xlsx())["text"]
        self.assertIn("【费用】", text_)
        self.assertIn("2026-10-01 | 260 | 高铁票", text_)

    def test_pdf_text_layer_and_page_count(self):
        result = self.parse("pdf", make_pdf())
        self.assertIn("Meeting minutes 2026 owner demo_emp", result["text"])
        self.assertEqual(result["pages"], 1)
        self.assertTrue(any("没有做图片文字识别" in w for w in result["warnings"]))

    def test_broken_pdf_is_a_clear_error(self):
        with self.assertRaises(InvalidInput):
            self.parse("pdf", b"%PDF-1.4 garbage")

    def test_pdf_page_cap(self):
        with patch.object(intake, "MAX_PDF_PAGES", 0), self.assertRaises(InvalidInput) as ctx:
            self.parse("pdf", make_pdf())
        self.assertIn("上限", ctx.exception.message)

    def test_email_headers_body_and_attachment_names(self):
        result = self.parse("eml", make_eml())
        self.assertIn("主题：关于下周发布的安排", result["text"])
        self.assertIn("发件人：boss@example.com", result["text"])
        self.assertIn("请 demo_emp 周五前提交方案。", result["text"])
        self.assertIn("【附件（未解析）】方案.zip", result["text"])
        self.assertEqual(result["attachments"], ["方案.zip"])

    def test_html_only_email_is_stripped_of_tags_and_scripts(self):
        text_ = self.parse("eml", make_eml(html_only=True))["text"]
        self.assertIn("请 demo_emp 周五前提交方案", text_)
        self.assertNotIn("<p>", text_)
        self.assertNotIn("alert", text_)

    def test_text_encodings(self):
        self.assertEqual(self.parse("txt", "会议纪要".encode("gb18030"))["text"], "会议纪要")
        self.assertEqual(self.parse("csv", "a,b\n1,2".encode("utf-8"))["text"], "a,b\n1,2")

    def test_temp_files_are_removed(self):
        before = {n for n in os.listdir(tempfile.gettempdir()) if n.startswith("intake_")}
        self.parse("docx", make_docx())
        self.assertEqual({n for n in os.listdir(tempfile.gettempdir()) if n.startswith("intake_")}, before)

    def test_clean_removes_control_characters_and_blank_runs(self):
        self.assertEqual(intake._clean("a\x00b\r\n\r\n\r\n\r\nc\x07 \n"), "ab\n\nc")


class ExtractTest(unittest.TestCase):
    """不连数据库：成员校验换成替身。"""

    def run_extract(self, name, content, sensitivity="internal", vision=None):
        with patch("service.automation_work_service.authorize", AsyncMock()), \
                patch.object(intake, "_vision_model", return_value=vision), \
                patch("service.audit_service.record_async", AsyncMock()) as audit:
            result = run_async(intake.extract_text(None, 1, 2, name, content, sensitivity))
        return result, audit

    def test_returns_text_summary_and_audits_without_content(self):
        result, audit = self.run_extract("纪要.docx", make_docx())
        self.assertEqual((result["file_type"], result["label"], result["file_name"]), ("docx", "Word", "纪要.docx"))
        self.assertEqual(result["chars"], len(result["text"]))
        detail = audit.call_args.kwargs["detail"]
        self.assertEqual(set(detail), {"type", "size", "chars", "ocr"})   # 审计里只有类型/大小，没有内容或文件名

    def test_rejections(self):
        cases = (("a.exe", b"MZ", "不支持"), ("a.txt", b"", "空的"), ("a.pdf", b"plain text", "不符"),
                 ("a.txt", b"x" * (intake.MAX_BYTES + 1), "上限"))
        for name, content, word in cases:
            with self.subTest(name=name), self.assertRaises(InvalidInput) as ctx:
                self.run_extract(name, content)
            self.assertIn(word, ctx.exception.message)

    def test_restricted_material_is_never_imported(self):
        with self.assertRaises(InvalidInput):
            self.run_extract("a.txt", b"hello world", "restricted")
        with self.assertRaises(InvalidInput):
            self.run_extract("a.txt", b"hello world", "secret")

    def test_image_needs_a_vision_model_that_the_sensitivity_allows(self):
        with self.assertRaises(InvalidInput) as ctx:
            self.run_extract("a.png", make_png())
        self.assertIn("配置支持视觉的模型", ctx.exception.message)
        vision = ("gpt-4o", "key", None)
        with self.assertRaises(InvalidInput) as ctx:
            self.run_extract("a.png", make_png(), "confidential", vision)   # gpt-4o 不在机密材料的受信任模型名单里
        self.assertIn("密级不允许", ctx.exception.message)

    def test_image_ocr_goes_through_the_configured_vision_model(self):
        vision = ("gpt-4o", "key", None)
        with patch("service.llm.vision_ocr.ocr_image", return_value="发票号 G001 金额 260 元"),                 patch("service.llm.vision_ocr.resolve_vision_model", return_value=vision):
            result, _ = self.run_extract("发票.png", make_png(), "internal", ("gpt-4o", "key", None))
        self.assertIn("G001", result["text"])

    def test_scanned_pdf_without_vision_model_explains_what_to_do(self):
        with self.assertRaises(InvalidInput) as ctx:
            self.run_extract("扫描件.pdf", make_pdf(scanned=True))
        self.assertIn("没有从文件里提取到文字", ctx.exception.message)

    def test_confidential_pdf_skips_ocr_and_says_so(self):
        result, _ = self.run_extract("a.pdf", make_pdf(), "confidential", ("gpt-4o", "key", None))
        self.assertTrue(any("没有做图片文字识别" in w for w in result["warnings"]))

    def test_long_text_is_cut_with_a_warning(self):
        with patch.object(intake, "MAX_TEXT_CHARS", 50):
            result, _ = self.run_extract("a.txt", ("行" * 80).encode("utf-8"))
        self.assertEqual(len(result["text"]), 50)
        self.assertTrue(any("只保留了前面" in w for w in result["warnings"]))


@unittest.skipUnless(_AVAILABLE, _WHY)
class ImportRouteTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.db = SessionLocal()
        cls.client = rc.make_client()
        cls.owner = rc.create_user("imp-owner")
        cls.member = rc.create_user("imp-mem")
        cls.stranger = rc.create_user("imp-out")
        cls.org = _create_org(cls.db, "imp-org-" + uuid.uuid4().hex[:6], cls.owner["id"])
        cls.team = _create_team(cls.db, cls.org, "imp-team", cls.owner["id"])
        cls.db.commit()
        _add_org_member(cls.db, cls.org, cls.member["id"], "member")
        _add_team_member(cls.db, cls.team, cls.member["id"], "member")

    @classmethod
    def tearDownClass(cls):
        rc.cleanup()
        cls.db.execute(text("DELETE FROM team_members WHERE team_id=:t"), {"t": cls.team})
        cls.db.execute(text("DELETE FROM teams WHERE id=:t"), {"t": cls.team})
        cls.db.execute(text("DELETE FROM organization_members WHERE organization_id=:o"), {"o": cls.org})
        cls.db.execute(text("DELETE FROM organizations WHERE id=:o"), {"o": cls.org})
        cls.db.commit()
        cls.db.close()

    def post(self, user, name, content, **form):
        data = {"team_id": str(self.team), **form}
        return self.client.post("/enterprise/automation/import", headers=user["headers"], data=data,
                                files={"file": (name, content, "application/octet-stream")})

    def test_member_imports_a_word_file(self):
        response = self.post(self.member, "纪要.docx", make_docx())
        self.assertEqual(response.status_code, 200, response.text)
        body = response.json()
        self.assertIn("例会纪要", body["text"])
        self.assertEqual(body["label"], "Word")

    def test_non_member_is_rejected_before_any_parsing(self):
        with patch.object(intake, "_parse_sync") as parse:
            response = self.post(self.stranger, "纪要.docx", make_docx())
        self.assertEqual(response.status_code, 403)
        parse.assert_not_called()

    def test_bad_input_is_a_400_with_a_readable_message(self):
        response = self.post(self.member, "a.exe", b"MZ....")
        self.assertEqual(response.status_code, 400)
        self.assertIn("不支持的文件类型", response.json()["detail"])
        self.assertEqual(self.post(self.member, "a.txt", b"hello world", sensitivity="restricted").status_code, 400)
        self.assertEqual(self.post(self.member, "a.txt", b"hello world", sensitivity="bogus").status_code, 422)

    def test_requires_login(self):
        response = self.client.post("/enterprise/automation/import", data={"team_id": str(self.team)},
                                    files={"file": ("a.txt", b"hello world", "text/plain")})
        self.assertEqual(response.status_code, 401)

    def test_extracted_text_feeds_the_normal_flow_without_being_stored(self):
        self.post(self.member, "a.txt", "例会纪要".encode("utf-8"))
        count = self.db.execute(text("SELECT COUNT(*) FROM automation_work WHERE team_id=:t"), {"t": self.team}).scalar()
        self.db.commit()
        self.assertEqual(count, 0)   # 导入只提取文字，不会生成工作成果、不保存文件


if __name__ == "__main__":
    unittest.main()
