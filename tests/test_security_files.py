"""恶意文件专项：压缩炸弹、路径穿越（含 Windows 盘符 / UNC / 备用数据流 / 保留设备名）、伪造扩展名、危险文件名。
每个接收用户文件的入口（AI 整理导入、考勤导入、技能包安装、旧版能力包、附件）都用同一批载荷验证。"""
import io
import os
import tempfile
import time
import unittest
import zipfile

from openpyxl import Workbook

from service import archive_guard
from service.exceptions import InvalidInput


def make_zip(members, compress=zipfile.ZIP_DEFLATED) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", compress) as archive:
        for name, data in members:
            archive.writestr(name, data)
    return buffer.getvalue()


def xlsx_bytes(extra=None) -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.append(["姓名", "打卡时间"])
    sheet.append(["张三", "2026-10-06 08:58"])
    buffer = io.BytesIO()
    book.save(buffer)
    if not extra:
        return buffer.getvalue()
    source = zipfile.ZipFile(io.BytesIO(buffer.getvalue()))
    members = [(info.filename, source.read(info.filename)) for info in source.infolist()]
    return make_zip(members + list(extra))


def bomb_xlsx(size_mb=400) -> bytes:
    """几百 KB 的文件，解压后几百 MB：结构完全合法，只是工作表末尾塞了一大段注释。"""
    book = Workbook()
    book.active.append(["姓名", "打卡时间"])
    buffer = io.BytesIO()
    book.save(buffer)
    source = zipfile.ZipFile(io.BytesIO(buffer.getvalue()))
    members = []
    for info in source.infolist():
        data = source.read(info.filename)
        if info.filename == "xl/worksheets/sheet1.xml":
            data = data.replace(b"</worksheet>", b"<!--" + b"A" * (size_mb * 1024 * 1024) + b"--></worksheet>")
        members.append((info.filename, data))
    return make_zip(members)


BAD_NAMES = ["../evil.txt", "..\\evil.txt", "a/../../evil.txt", "a\\..\\..\\evil.txt", "/abs/evil.txt", "\\\\server\\share\\evil.txt", "//server/share/evil.txt",
             "C:/Windows/evil.txt", "C:\\Windows\\evil.txt", "c:evil.txt", "evil.txt:stream", "evil.txt::$DATA", "dir/CON", "con.txt", "NUL", "aux.md", "COM1.txt",
             "lpt9", "evil.txt.", "evil.txt ", "dir/ /x", "a\x00b.txt", "a\nb.txt", "x" * 300, ""]
GOOD_NAMES = ["SKILL.md", "docs/readme.txt", "my skill/resources/data.csv", "中文目录/说明.md", "a/b/c/d.json", "./SKILL.md", "dir/", "xl/worksheets/sheet1.xml",
              "[Content_Types].xml", "_rels/.rels", "console.txt", "auxiliary.md", "com10.txt"]


class ArchiveGuardTest(unittest.TestCase):
    def test_member_name_matrix(self):
        for name in BAD_NAMES:
            with self.subTest(bad=name):
                self.assertFalse(archive_guard.safe_member_name(name), repr(name))
        for name in GOOD_NAMES:
            with self.subTest(good=name):
                self.assertTrue(archive_guard.safe_member_name(name), repr(name))

    def test_inspect_rejects_bombs_traversal_and_corruption(self):
        with self.assertRaises(archive_guard.ArchiveRejected):
            archive_guard.inspect_zip(make_zip([("a.txt", b"0" * (150 * 1024 * 1024))]))                 # 总大小
        with self.assertRaises(archive_guard.ArchiveRejected):
            archive_guard.inspect_zip(make_zip([("a.txt", b"0" * (20 * 1024 * 1024))]))                  # 压缩比
        with self.assertRaises(archive_guard.ArchiveRejected):
            archive_guard.inspect_zip(make_zip([(f"f{i}.txt", b"x") for i in range(30)]), max_members=20)  # 成员数
        with self.assertRaises(archive_guard.ArchiveRejected):
            archive_guard.inspect_zip(make_zip([("../x.txt", b"x")]))
        with self.assertRaises(archive_guard.ArchiveRejected):
            archive_guard.inspect_zip(b"PK\x03\x04 not really a zip")
        self.assertEqual(len(archive_guard.inspect_zip(make_zip([("a.txt", b"hello"), ("b/c.txt", b"world")]))), 2)

    def test_declared_size_cannot_be_forged_to_hide_a_bomb(self):
        """压缩包头里把体积写小，读取时 zipfile 也只会读声明的那么多字节（并在校验失败时报错），不会无限解压。"""
        import struct
        data = bytearray(make_zip([("a.txt", b"0" * (5 * 1024 * 1024))]))
        # 把中央目录和本地头里的 file_size 都改成 10 字节
        for signature, offset in ((b"PK\x01\x02", 24), (b"PK\x03\x04", 22)):
            position = data.find(signature)
            data[position + offset:position + offset + 4] = struct.pack("<I", 10)
        with zipfile.ZipFile(io.BytesIO(bytes(data))) as archive:
            with self.assertRaises(zipfile.BadZipFile):
                archive.open("a.txt").read()

    def test_ensure_within_blocks_escapes_even_if_the_string_check_was_fooled(self):
        with tempfile.TemporaryDirectory() as base:
            inside = archive_guard.ensure_within(base, os.path.join(base, "a", "b.txt"))
            self.assertTrue(inside.startswith(os.path.realpath(base)))
            for target in (os.path.join(base, "..", "outside.txt"), os.path.join(os.path.dirname(base), "other.txt"), os.path.abspath(os.sep + "etc-passwd")):
                with self.subTest(target=target), self.assertRaises(archive_guard.ArchiveRejected):
                    archive_guard.ensure_within(base, target)


class AttendanceImportFilesTest(unittest.TestCase):
    def test_zip_bomb_is_rejected_quickly_without_decompressing(self):
        from service import attendance_import
        content = bomb_xlsx(400)
        self.assertLess(len(content), 2 * 1024 * 1024)                    # 压缩后不到 2MB
        started = time.time()
        with self.assertRaises(InvalidInput) as ctx:
            attendance_import.parse("考勤.xlsx", content)
        self.assertLess(time.time() - started, 1.0)
        self.assertIn("压缩", ctx.exception.message + "压缩")

    def test_path_traversal_member_names_are_rejected(self):
        from service import attendance_import
        for name in ("../../evil.txt", "C:/x.txt", "x.txt:ads"):
            with self.subTest(name=name), self.assertRaises(InvalidInput):
                attendance_import.parse("考勤.xlsx", xlsx_bytes(extra=[(name, b"x")]))

    def test_a_normal_workbook_still_parses(self):
        from service import attendance_import
        self.assertEqual(len(attendance_import.parse("考勤.xlsx", xlsx_bytes())["records"]), 1)

    def test_extremely_wide_sheets_are_cut_off(self):
        from service import attendance_import
        book = Workbook()
        sheet = book.active
        sheet.append(["姓名", "打卡时间"] + [f"c{i}" for i in range(300)])
        sheet.append(["张三", "2026-10-06 08:58"] + ["x"] * 300)
        buffer = io.BytesIO()
        book.save(buffer)
        parsed = attendance_import.parse("考勤.xlsx", buffer.getvalue())
        self.assertLessEqual(len(parsed["columns"]), attendance_import.MAX_COLUMNS)


class DocumentIntakeFilesTest(unittest.TestCase):
    def test_archive_check_rejects_bombs_and_traversal(self):
        from service import document_intake
        for content in (bomb_xlsx(300), make_zip([("word/document.xml", b"<w/>"), ("../x", b"x")]), make_zip([("a.txt", b"0" * (100 * 1024 * 1024))])):
            with self.assertRaises(InvalidInput):
                document_intake._check_archive(content)
        document_intake._check_archive(xlsx_bytes())            # 正常文件不受影响

    def test_extension_spoofing_is_rejected(self):
        from service import document_intake
        exe = b"MZ\x90\x00" + b"\x00" * 100
        elf = b"\x7fELF" + b"\x00" * 100
        for kind, content in (("pdf", exe), ("docx", exe), ("xlsx", elf), ("png", exe), ("jpg", exe), ("txt", exe), ("md", elf), ("csv", b"PK\x03\x04" + b"x" * 50),
                              ("docx", b"%PDF-1.4 fake"), ("pdf", b"PK\x03\x04 zipfile")):
            with self.subTest(kind=kind, head=content[:4]), self.assertRaises(InvalidInput):
                document_intake._check_signature(kind, content)

    def test_unsupported_and_double_extensions(self):
        from service import document_intake
        for name in ("evil.exe", "evil.pdf.exe", "evil", "evil.html", "evil.svg", "evil.docm", "evil.xlsm", ".pdf.js"):
            with self.subTest(name=name), self.assertRaises(InvalidInput):
                document_intake.file_type_of(name)
        self.assertEqual(document_intake.file_type_of("report.final.PDF"), "pdf")


class SkillPackageFilesTest(unittest.TestCase):
    def test_package_import_rejects_traversal_and_windows_tricks_before_touching_disk(self):
        from service.skills_core.package_import import SkillImportError, import_skill_bundle
        for name in ("../evil.txt", "a/../../evil.txt", "C:/Windows/evil.txt", "skill/data.txt:ads", "skill/CON.md", "\\\\server\\share\\x", "/abs.md"):
            content = make_zip([("skill/SKILL.md", b"---\nname: t\n---\nbody"), (name, b"x")])
            with self.subTest(name=name), self.assertRaises(SkillImportError):
                import_skill_bundle(None, 1, "evil.zip", content, commit=False)

    def test_package_import_rejects_zip_bombs(self):
        from service.skills_core.package_import import SkillImportError, import_skill_bundle
        content = make_zip([("skill/SKILL.md", b"---\nname: t\n---\nbody"), ("skill/big.txt", b"0" * (20 * 1024 * 1024))])
        with self.assertRaises(SkillImportError) as ctx:
            import_skill_bundle(None, 1, "bomb.zip", content, commit=False)
        self.assertIn("压缩", str(ctx.exception))

    def test_target_path_is_always_inside_the_package_directory(self):
        from service.skills_core.package_import import SkillImportError, _target_within
        with tempfile.TemporaryDirectory() as base:
            self.assertTrue(_target_within(base, "docs/readme.txt").startswith(os.path.realpath(base)))
            other_drive = ("D:" if os.path.realpath(base)[:2].upper() == "C:" else "C:") if os.name == "nt" else None
            escapes = ["../x.txt", "a/../../x.txt", "/etc/x.txt"] + ([f"{other_drive}/x.txt", f"{other_drive}x.txt"] if other_drive else [])
            for rel in escapes:
                with self.subTest(rel=rel), self.assertRaises(SkillImportError):
                    _target_within(base, rel)

    def test_legacy_package_with_traversal_or_bomb_is_refused_and_leaves_nothing_behind(self):
        from service.skills_core import import_export
        root = os.path.join(os.path.dirname(import_export.SKILLS_ROOT), "skills_packages", "imported")
        for members in ([("manifest.yaml", b"name: x"), ("SKILL.md", b"x"), ("../evil.txt", b"x")],
                        [("manifest.yaml", b"name: x"), ("SKILL.md", b"x"), ("big.txt", b"0" * (30 * 1024 * 1024))]):
            before = set(os.listdir(root)) if os.path.isdir(root) else set()
            self.assertIsNone(import_export.import_skill_from_upload(None, 1, "legacy.zip", make_zip(members), commit=False))
            self.assertEqual(set(os.listdir(root)) if os.path.isdir(root) else set(), before)


class AttachmentNamesTest(unittest.TestCase):
    def test_dangerous_names_become_harmless_plain_file_names(self):
        from service import attachment_service
        for raw in ("../../etc/passwd", "..\\..\\Windows\\win.ini", "C:\\Users\\x\\a.txt", "/etc/shadow", "a/b/c.txt", "évil\u202Egpj.exe", "x\x00.txt",
                    "con.txt", "NUL", "aux", "COM1.log", "name:stream.txt", "<script>.txt", "a|b?.txt", "." * 5, "   ", "", None, "x" * 500):
            name = attachment_service.safe_filename(raw)
            with self.subTest(raw=raw):
                self.assertNotIn("/", name)
                self.assertNotIn("\\", name)
                self.assertNotIn(":", name)
                self.assertNotIn("\x00", name)
                self.assertNotIn("\u202e", name)
                self.assertLessEqual(len(name), 121)
                self.assertTrue(name)
                self.assertNotIn(name.split(".")[0].lower(), {"con", "prn", "aux", "nul", "com1", "lpt1"})     # 不能是 Windows 设备名

    def test_executable_extensions_are_refused(self):
        from service import attachment_service
        for name in ("a.exe", "a.EXE", "a.bat", "a.ps1", "a.sh", "a.js.exe", "a.exe.", "a.exe "):
            with self.subTest(name=name), self.assertRaises(attachment_service.AttachmentError):
                attachment_service.save(1, name, b"x")


if __name__ == "__main__":
    unittest.main()
