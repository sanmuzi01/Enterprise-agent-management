"""
RAG编排服务：把文档解析→切分→嵌入→向量存储→DB存储 全部串起来
对外提供3个高级接口：
  upload_and_index()              上传文档并入库
  search()                        根据问题检索知识片段
  delete_knowledge_completely()   彻底删除文档（向量库+chunks+knowledge记录）
这一层在架构中的位置
  路由层 (FasdtApi/knowledge.py)
        ↓
  rag_service.py (业务编排 ← 你在这里)
        ↓
  embedding_service + vector_store_service + knowledge_dao + knowledge_chunk_dao
        ↓
  智谱API + ChromaDB + MySQL
路由层不该关心"怎么切分文档""怎么调嵌入API"，
这些业务逻辑封装在这里，路由层只调 upload_and_index() / search()。
"""
import base64
import os
import re
import uuid
from typing import List,Dict,Any
from dotenv import load_dotenv
from service.rag.rerank.factory import RerankFactory
from models.knowledge_chunk_dao import create_chunks_batch, get_chunks_by_vector_ids, delete_chunks_by_knowledge
from utils.logger_handler import get_logger
from utils.path_tool import get_abs_path
from utils.upload_limits import StoredUpload
# 导入 RAG 底层两层能力
from service.rag.embedding_service import embed_texts,embed_query
from service.rag.embedding_service import aembed_query_async
from service.rag.vector_store_service import (add_vectors,search_similar,delete_vectors_by_knowledge,space_collection_key)
# 导入 DAO（4层架构：Service层只调DAO，不直接碰 ORM）
from models.knowledge_dao import (
    create_knowledge,get_knowledge_by_id,get_knowledge_by_ids,
    delete_knowledge as dao_delete_knowledge,
    update_knowledge_status)

load_dotenv()
# 上传的原始文档保存在哪（磁盘）
KNOWLEDGE_FILE_PATH = os.getenv("KNOWLEDGE_FILE_PATH","./knowledge_files")
# 切块参数
CHUNK_SIZE = 500        # 每块约500字符（用户没自选时的默认）
CHUNK_OVERLAP = 50      # 相邻块重叠50字符（防止一句话被切断）
CHUNK_SIZE_MIN = 120    # 用户可选切块大小下限
CHUNK_SIZE_MAX = 2000   # 上限


def _effective_chunk_params(knowledge) -> tuple:
    """按文档的 chunk_size（用户自选）解析出 (chunk_size, overlap)，越界夹紧，None 用默认。

    overlap 固定 CHUNK_OVERLAP，仅在 chunk 很小时按 size//4 缩小（保证 overlap < size）。
    """
    raw = getattr(knowledge, "chunk_size", None)
    size = CHUNK_SIZE if not raw else max(CHUNK_SIZE_MIN, min(int(raw), CHUNK_SIZE_MAX))
    overlap = min(CHUNK_OVERLAP, size // 4)
    return size, overlap


def clamp_chunk_size(value) -> "int | None":
    """路由层入口校验：把用户传的 chunk_size 夹到合法区间；空/非法 → None（用默认）。"""
    try:
        n = int(value)
    except (TypeError, ValueError):
        return None
    return max(CHUNK_SIZE_MIN, min(n, CHUNK_SIZE_MAX))
RAG_RERANK_ENABLED = os.getenv("RAG_RERANK_ENABLED", "0").lower() in {"1", "true", "yes", "on"}
RAG_RERANK_MODEL = os.getenv("RAG_RERANK_MODEL", "BAAI/bge-reranker-v2-m3")
logger = get_logger("rag_service")
# ========== 获取 Rerank Client ==========
_rerank_client = None
def _get_rerank_client():
    """获取Rerank客户端（本地BGE模型，不需要Key）"""
    global _rerank_client
    if not RAG_RERANK_ENABLED:
        logger.info("Rerank未启用，跳过重排序")
        return None
    if _rerank_client is not None:
        return _rerank_client
    try:
        logger.info(f"Rerank已启用，开始加载模型: {RAG_RERANK_MODEL}")
        _rerank_client = RerankFactory.create(RAG_RERANK_MODEL)
        return _rerank_client
    except Exception as e:
        logger.warning(f"创建Reranker失败: {e}，将跳过重排序")
        return None
# ========== 文档解析（不同文件类型用不同库） ==========

MAX_OCR_PAGES_PER_DOCUMENT = int(os.getenv("RAG_MAX_OCR_PAGES_PER_DOCUMENT", "30"))


def _parse_pdf(file_path: str, db=None, user_id: int = None) -> str:
    """解析PDF成纯文本。每页前插一个 [第N页] 标记——纯文本层面最简单的定位手段：
    某个 chunk 切到哪页，标记会随着切块一起留在 chunk 内容里，引用展示时用户
    一眼就知道这段话来自原 PDF 第几页，不用额外的页码字段/schema 改动。

    扫描件/纯图片页提取不出文字时，如果调用方传了 db/user_id 且用户配置了视觉模型
    （GLM-4V / GPT-4o），会尝试用视觉模型 OCR 出文字，标成 [第N页·OCR识别]；
    没配视觉模型、或 OCR 失败，就沿用老行为跳过这页，不让整份文档处理失败。
    单份文档最多 OCR MAX_OCR_PAGES_PER_DOCUMENT 页（默认30），避免超大扫描件
    在后台任务里跑出几十上百次模型调用。
    """
    from pypdf import PdfReader

    reader = PdfReader(file_path)
    parts = []
    vision_model = None
    vision_resolved = False
    render_doc = None
    ocr_used = 0
    try:
        for page_no, page in enumerate(reader.pages, start=1):
            page_text = (page.extract_text() or "").strip()
            if page_text:
                parts.append(f"[第{page_no}页]\n{page_text}")
                continue
            # 扫描件/纯图片页：没有 db/user_id（比如未来的非文档场景调用）就直接跳过
            if db is None or user_id is None or ocr_used >= MAX_OCR_PAGES_PER_DOCUMENT:
                continue
            if not vision_resolved:
                from service.llm.vision_ocr import resolve_vision_model
                vision_model = resolve_vision_model(db, user_id)
                vision_resolved = True
            if not vision_model:
                continue
            try:
                if render_doc is None:
                    import pymupdf  # 未安装时跳过 OCR，不影响其它页的正常解析
                    render_doc = pymupdf.open(file_path)
                pixmap = render_doc[page_no - 1].get_pixmap(dpi=200)
                image_b64 = base64.b64encode(pixmap.tobytes("png")).decode("ascii")
                from service.llm.vision_ocr import ocr_image
                model_name, api_key, api_url = vision_model
                ocr_text = ocr_image(model_name, api_key, api_url, image_b64)
                ocr_used += 1
                if ocr_text:
                    parts.append(f"[第{page_no}页·OCR识别]\n{ocr_text}")
            except ImportError:
                logger.warning("未安装 pymupdf，扫描件页无法 OCR，已跳过")
                vision_model = None  # 后续页不用再试，省得重复触发同一个 ImportError
            except Exception as e:  # noqa: BLE001 - OCR 失败只影响这一页，不影响整份文档
                logger.warning(f"OCR 识别第{page_no}页失败，跳过: {e}")
    finally:
        if render_doc is not None:
            render_doc.close()
    return "\n\n".join(parts)


def _parse_xlsx(file_path: str, db=None, user_id: int = None) -> str:
    """解析 Excel 成纯文本：每个工作表一段，每行用 | 分隔单元格，跳过完全空白的行。"""
    from openpyxl import load_workbook

    workbook = load_workbook(file_path, read_only=True, data_only=True)
    try:
        parts = []
        for sheet in workbook.worksheets:
            rows = []
            for row in sheet.iter_rows(values_only=True):
                cells = ["" if v is None else str(v).strip() for v in row]
                if any(cells):
                    rows.append(" | ".join(cells))
            if rows:
                parts.append(f"【{sheet.title}】\n" + "\n".join(rows))
        return "\n\n".join(parts)
    finally:
        workbook.close()


def _parse_image(file_path: str, db=None, user_id: int = None) -> str:
    """纯图片文件：整张图交给视觉模型 OCR。

    和扫描件 PDF 页不同——图片上传的全部意义就是里面的文字，没有"提取不到就跳过"
    这个退路，没配视觉模型或识别失败都应该是明确的错误，而不是悄悄产出一份空文档。
    """
    if db is None or user_id is None:
        raise ValueError("图片文件需要 OCR 识别，当前调用场景不支持")

    from service.llm.vision_ocr import ocr_image, resolve_vision_model

    vision_model = resolve_vision_model(db, user_id)
    if not vision_model:
        raise ValueError(
            "识别图片文字需要先配置一个支持视觉的模型（智谱 GLM-4V 或 OpenAI GPT-4o/GPT-4o-mini），"
            "请到「模型连接」页添加后重试。"
        )
    ext = os.path.splitext(file_path)[1].lstrip(".").lower()
    mime_type = "image/jpeg" if ext in ("jpg", "jpeg") else f"image/{ext or 'png'}"
    with open(file_path, "rb") as f:
        image_b64 = base64.b64encode(f.read()).decode("ascii")
    model_name, api_key, api_url = vision_model
    return ocr_image(model_name, api_key, api_url, image_b64, mime_type=mime_type)




def _iter_docx_block_items(doc):
    """按文档里出现的先后顺序，依次产出段落(Paragraph)和表格(Table)。

    python-docx 没有直接提供"正文顺序遍历"的 API——doc.paragraphs 和 doc.tables
    是分开的两个列表，天然就丢了表格在文中的相对位置。这是 python-docx 官方文档
    推荐的标准写法：直接读 body 的 XML 子节点，按类型分发。
    """
    from docx.document import Document as _Document
    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    parent_elm = doc.element.body if isinstance(doc, _Document) else doc._element
    for child in parent_elm.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, doc)
        elif isinstance(child, CT_Tbl):
            yield Table(child, doc)


def _table_to_text(table) -> str:
    """表格转成分隔符文本：每行一条，单元格用 | 分开，整体裹一个【表格】标记。

    之前 _parse_docx 只读 doc.paragraphs，表格内容完全没进入过切块/向量化——
    产品说明书里常见的参数表、价目表全部丢失。转纯文本虽然丢掉了表格的视觉结构，
    但内容不再丢失，向量检索也能命中表格里的关键词。
    """
    rows = []
    for row in table.rows:
        cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
        if any(cells):
            rows.append(" | ".join(cells))
    if not rows:
        return ""
    return "【表格】\n" + "\n".join(rows)


def _parse_docx(file_path: str, db=None, user_id: int = None) -> str:
    """解析 Word(.docx) 成纯文本，段落和表格按文中原有顺序输出。"""
    import docx
    doc = docx.Document(file_path)
    parts = []
    for block in _iter_docx_block_items(doc):
        if hasattr(block, "rows"):  # Table
            table_text = _table_to_text(block)
            if table_text:
                parts.append(table_text)
        else:  # Paragraph
            if block.text.strip():
                parts.append(block.text)
    return "\n".join(parts)
def _parse_txt(file_path: str, db=None, user_id: int = None) -> str:
    """解析纯文本 / Markdown"""
    for encoding in ("utf-8", "utf-8-sig", "gb18030"):
        try:
            with open(file_path, "r", encoding=encoding) as f:
                return f.read()
        except UnicodeDecodeError:
            continue
    with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
        return f.read()
# 支持的文件类型 → 解析函数，单一事实来源。路由层的 ALLOWED_TYPES 应该从
# SUPPORTED_FILE_TYPES 派生，不要各自维护一份硬编码列表（否则新增类型要改两处，
# 之前 FasdtApi/knowledge.py 和 service/knowledge_space/document_service.py
# 就各自维护了一份一模一样的 {"txt","md","pdf","docx"}，容易改漏）。
_PARSERS = {
    "pdf": _parse_pdf,
    "docx": _parse_docx,
    "txt": _parse_txt,
    "md": _parse_txt,  # markdown 和纯文本解析方式一样
    "xlsx": _parse_xlsx,
    "jpg": _parse_image,
    "jpeg": _parse_image,
    "png": _parse_image,
}
SUPPORTED_FILE_TYPES = sorted(_PARSERS.keys())


def parse_document(file_path: str, file_type: str, db=None, user_id: int = None) -> str:
    """根据文件类型选择解析器。
      新增文件类型时只加一行 key-value，不用改 if 结构。
      而且抛错时能直接给出"不支持的类型"，不用写默认else。
      db/user_id 是可选的——只有 PDF 扫描页 OCR 和纯图片 OCR 需要它们去查用户配置的
      视觉模型，其它解析器直接忽略。
      """
    parser = _PARSERS.get(file_type)
    if parser is None:
        raise ValueError(f"不支持的文件类型: {file_type}，支持: {SUPPORTED_FILE_TYPES}")
    return parser(file_path, db=db, user_id=user_id)
# ========== 文本切分 ==========
def _split_fixed_window(text: str, chunk_size: int, overlap: int) -> List[str]:
    """定长滑窗切分（旧算法）。只在单个自然段/表格本身就超过 chunk_size 时兜底用，
    保证再长的一块内容也能落地，不会因为切不动而卡死。"""
    chunks = []
    start = 0
    while start < len(text):
        end = start + chunk_size
        chunk = text[start:end]
        chunks.append(chunk.strip())
        start = end - overlap
    return chunks


def split_text(
        text:str,chunk_size:int = CHUNK_SIZE,
        overlap:int = CHUNK_OVERLAP
)->List[str]:
    """按段落切分：优先在空行处断开，尽量不把一句话/一段表格从中间切断，
    只有单个自然段本身就超过 chunk_size 才退回定长滑窗切。
    之前是不管内容结构、纯按字符数切，经常把一句话切成两半分到不同 chunk 里，
    检索命中其中一半时上下文残缺，也更容易在编号、型号这类需要完整读到的地方出问题。
    """
    if not text or not text.strip():
        return []

    # 空行是最自然的段落边界；PDF 页标记 [第N页]、表格标记【表格】也都是独立成段的，
    # 天然不会被硬切断。
    blocks = [b.strip() for b in re.split(r"\n\s*\n", text) if b.strip()]
    if not blocks:
        blocks = [text.strip()]

    chunks: List[str] = []
    current = ""
    for block in blocks:
        candidate = f"{current}\n\n{block}" if current else block
        if len(candidate) <= chunk_size:
            current = candidate
            continue
        if current:
            chunks.append(current)
        if len(block) <= chunk_size:
            current = block
        else:
            # 罕见情况：单个自然段/表格本身就超过 chunk_size，退回定长滑窗切
            sub_chunks = _split_fixed_window(block, chunk_size, overlap)
            if sub_chunks:
                chunks.extend(sub_chunks[:-1])
                current = sub_chunks[-1]
            else:
                current = ""
    if current:
        chunks.append(current)

    # 相邻块之间带一点上一块的尾巴，缓解"关键信息刚好卡在切分点"的问题
    if overlap > 0 and len(chunks) > 1:
        with_overlap = [chunks[0]]
        for i in range(1, len(chunks)):
            prev_tail = chunks[i - 1][-overlap:]
            with_overlap.append(f"{prev_tail}\n{chunks[i]}")
        chunks = with_overlap

    # 过滤掉太短的块（小于十个字符的多半是切分后的残留空白）
    chunks = [c.strip() for c in chunks if len(c.strip()) > 10]
    logger.info(f"文本切分完成（按段落）：{len(chunks)} 块，目标每块约 {chunk_size} 字")
    return chunks
# ========== 上传入库完整流程（核心） ==========
def _vector_key(knowledge):
    """向量集合键：优先 space_<id>，否则回退旧 agent_id。"""
    sid = getattr(knowledge, "space_id", None)
    return space_collection_key(sid) if sid else knowledge.agent_id


def upload_and_index(
        db,user_id:int,agent_id,file_name:str,file_content:bytes,file_type:str
)->Dict[str,Any]:#返回字典
    knowledge = prepare_upload(db, user_id, agent_id, file_name, file_content, file_type)
    return index_existing_knowledge(db, user_id, agent_id, knowledge.id)


def prepare_upload(db, user_id: int, agent_id, file_name: str,
                   file_content: bytes, file_type: str, *, space_id: int = None,
                   category: str = None, tags_json: str = None, version: str = None,
                   source_type: str = "upload", source_url: str = None,
                   chunk_size: int = None):
    """保存原文件并创建 pending 文档记录，不执行耗时入库。"""
    file_dir = get_abs_path(KNOWLEDGE_FILE_PATH)
    os.makedirs(file_dir,exist_ok=True)#创建文件目录，如果文件夹存在则不报错
    safe_name = f"{uuid.uuid4().hex}_{file_name}"
    file_path = os.path.join(file_dir, safe_name)#把文件夹路径和文件名拼接成完整文件路径。
    if isinstance(file_content, StoredUpload):
        # 上传接口已经把内容流式写进了临时文件：直接移到最终位置（同一个文件系统上的 rename），不再经过内存
        file_size = file_content.size
        file_content.move_to(file_path)
    else:
        with open(file_path,"wb") as f:
            f.write(file_content)#文本内容打开写入二进制文件真是内容
        file_size = len(file_content)
    # ---------- 第2步：DB 建 knowledge 记录（status=pending） ----------
    knowledge = create_knowledge(
        db, user_id=user_id, agent_id=agent_id, file_name=file_name,
        file_path=file_path, file_type=file_type, file_size=file_size,
        space_id=space_id, category=category, tags_json=tags_json, version=version,
        source_type=source_type, source_url=source_url,
        chunk_size=clamp_chunk_size(chunk_size),
    )
    logger.info(f"创建待入库文档: knowledge_id={knowledge.id}, file={file_name}")
    return knowledge


def index_existing_knowledge(db, user_id: int, agent_id: int, knowledge_id: int) -> Dict[str, Any]:
    """解析已有文档并写入 chunks + 向量。"""
    knowledge = get_knowledge_by_id(db, knowledge_id)
    if not knowledge or knowledge.user_id != user_id or (agent_id is not None and knowledge.agent_id != agent_id):
        raise ValueError("文档不存在或无权限")
    if not os.path.exists(knowledge.file_path):
        update_knowledge_status(db, knowledge, "failed", error_msg="原始文件已丢失，请重新上传这份资料。")
        raise ValueError("原始文件已丢失，请重新上传这份资料。")
    try:
        # ---------- 第3步：更新状态 processing ----------
        update_knowledge_status(db,knowledge,"processing")#数据库，知识库文件记录，状态
        db.flush()
        # ---------- 第4步：解析文档 ----------
        text = parse_document(knowledge.file_path, knowledge.file_type, db=db, user_id=user_id)
        if not text.strip():
            raise ValueError("没读到文字内容——文档太短、全是空白/表格图片，或者是扫描件但没有配置视觉模型（智谱 GLM-4V / OpenAI GPT-4o）做 OCR 识别。")
        # ---------- 第5步：切分（按文档自选 chunk_size，没选用默认） ----------
        _cs, _ov = _effective_chunk_params(knowledge)
        chunks_text = split_text(text, chunk_size=_cs, overlap=_ov)
        if not chunks_text:
            raise ValueError("文档里几乎没有有效文字（可能太短，或全是空白 / 表格图片）。")
        # ---------- 第6步：批量嵌入（调智谱API，可能耗时） ----------
        vectors = embed_texts(db,user_id,chunks_text)
        # ---------- 第7步：存向量到 ChromaDB ----------
        # 构造有意义的 vector_id：k{knowledge_id}_c{chunk_index}
        vector_ids = [f"k{knowledge.id}_c{i}" for i in range(len(chunks_text))]#每条向量的唯一ID
        # metadata：存 knowledge_id 和 chunk_index，后面按文档删向量要用
        metadatas = [
            {"knowledge_id": knowledge.id, "chunk_index": i}
            for i in range(len(chunks_text))
        ]#给每个向量附加额外信息。
        add_vectors(
            agent_id = _vector_key(knowledge),vectors=vectors,
            ids=vector_ids,documents=chunks_text,metadatas=metadatas)
        # ---------- 第8步：批量存 chunks 到 MySQL ----------
        chunk_records=[
            {
                "knowledge_id": knowledge.id,
                "chunk_index": i,
                "content": chunks_text[i],
                "vector_id": vector_ids[i],
                "token_count": len(chunks_text[i]) // 2,  # 粗估：中文每2字符≈1token
            }
            for i in range(len(chunks_text))
        ]
        create_chunks_batch(db,chunk_records)
        # ---------- 第9步：更新状态 done，回填 chunk_count ----------
        update_knowledge_status(db,knowledge,"done",chunk_count=len(chunks_text))
        logger.info(f"文档入库成功: knowledge_id={knowledge.id}, chunks={len(chunks_text)}")
        return {
            "knowledge_id": knowledge.id,
            "chunk_count": len(chunks_text),
            "message": "上传并入库成功",
        }
    except Exception as e:
        # ---------- 任何一步失败：把状态标成 failed ----------
        update_knowledge_status(db,knowledge,"failed", error_msg=str(e))
        logger.error(f"文档入库失败: knowledge_id={knowledge.id}, error={e}")
        # 继续向上抛异常，让路由层返回给用户
        raise
    # ========== 检索流程（RAG的"查"） ==========


def reindex_knowledge(db, user_id: int, agent_id: int, knowledge_id: int) -> Dict[str, Any]:
    """重新解析已有文件并重建 chunks + 向量。

    顺序是关键：先把新内容全部生成好（解析 → 切块 → 向量化——这三步最容易失败：
    文件内容读不出来、切完没有有效文字、embedding API 报错/超限），确认新内容齐了
    再删旧数据、写新数据。这样"生成新内容"阶段任何一步失败，旧的 chunks/向量完全
    不受影响，文档仍能正常检索；不会像之前那样一上来就把旧数据删光，一旦后面任何
    一步失败，文档就变成没有任何内容、无法检索的空文档。
    """
    knowledge = get_knowledge_by_id(db, knowledge_id)
    if not knowledge or knowledge.user_id != user_id or (agent_id is not None and knowledge.agent_id != agent_id):
        raise ValueError("文档不存在或无权限")
    if not os.path.exists(knowledge.file_path):
        update_knowledge_status(db, knowledge, "failed", error_msg="原始文件已丢失，请重新上传这份资料。")
        raise ValueError("原始文件已丢失，请重新上传这份资料。")

    try:
        update_knowledge_status(db, knowledge, "processing")
        db.flush()

        # ---- 先把新内容全部生成好，旧数据这时候还原封不动 ----
        text = parse_document(knowledge.file_path, knowledge.file_type, db=db, user_id=user_id)
        if not text.strip():
            raise ValueError("没读到文字内容——文档太短、全是空白/表格图片，或者是扫描件但没有配置视觉模型（智谱 GLM-4V / OpenAI GPT-4o）做 OCR 识别。")
        _cs, _ov = _effective_chunk_params(knowledge)
        chunks_text = split_text(text, chunk_size=_cs, overlap=_ov)
        if not chunks_text:
            raise ValueError("文档里几乎没有有效文字（可能太短，或全是空白 / 表格图片）。")
        vectors = embed_texts(db, user_id, chunks_text)

        # ---- 新内容都拿到手了，才开始删旧、写新 ----
        try:
            delete_vectors_by_knowledge(_vector_key(knowledge), knowledge_id)
        except Exception as e:
            logger.warning(f"重建索引时删除旧向量失败，继续重建: {e}")
        delete_chunks_by_knowledge(db, knowledge_id)

        vector_ids = [f"k{knowledge.id}_c{i}" for i in range(len(chunks_text))]
        metadatas = [
            {"knowledge_id": knowledge.id, "chunk_index": i}
            for i in range(len(chunks_text))
        ]
        add_vectors(
            agent_id=_vector_key(knowledge),
            vectors=vectors,
            ids=vector_ids,
            documents=chunks_text,
            metadatas=metadatas,
        )
        chunk_records = [
            {
                "knowledge_id": knowledge.id,
                "chunk_index": i,
                "content": chunks_text[i],
                "vector_id": vector_ids[i],
                "token_count": len(chunks_text[i]) // 2,
            }
            for i in range(len(chunks_text))
        ]
        create_chunks_batch(db, chunk_records)
        update_knowledge_status(db, knowledge, "done", chunk_count=len(chunks_text))
        logger.info(f"文档重建索引成功: knowledge_id={knowledge.id}, chunks={len(chunks_text)}")
        return {
            "knowledge_id": knowledge.id,
            "chunk_count": len(chunks_text),
            "message": "重新入库成功",
        }
    except Exception as e:
        update_knowledge_status(db, knowledge, "failed", error_msg=str(e))
        logger.error(f"文档重建索引失败: knowledge_id={knowledge.id}, error={e}")
        raise


def search(
        db,user_id:int,agent_id:int ,query:str,top_k:int = 5, knowledge_id:int = None
)->List[Dict[str,Any]]:
    """完整检索（4步）：
        用户问题向量化,向量库相似度检索（拿到 vector_id 列表）
        反查 MySQL 拿完整 chunk 内容 + 按相似度排序,Rerank重排序（如果Reranker可用）
       """
    # 第1步：问题向量化
    if not query or not query.strip():
        raise ValueError("检索关键词不能为空")
    query_vector = embed_query(db,user_id,query)
    return _build_search_results(db, agent_id, query, top_k, query_vector, knowledge_id)


def _build_search_results(
        db, agent_id: int, query: str, top_k: int,
        query_vector: List[float], knowledge_id: int = None,
) -> List[Dict[str, Any]]:
    """根据已经生成好的查询向量完成向量检索、DB 反查和可选重排。"""

    if not query_vector:
        return []
    logger.info(f"开始检索: agent={agent_id}, query='{query[:50]}...', top_k={top_k}")
    retrieve_count = max(10, top_k * 2) if RAG_RERANK_ENABLED else top_k
    where = {"knowledge_id": knowledge_id} if knowledge_id is not None else None
    results = search_similar(agent_id, query_vector, top_k=retrieve_count, where=where)
    if not results:
        logger.info(f"向量检索无命中: agent={agent_id}, query='{query[:50]}...'")
        return []
    logger.info(f"向量检索完成: agent={agent_id}, 命中{len(results)}条")

    vector_ids = [r["id"] for r in results]
    chunks = get_chunks_by_vector_ids(db, vector_ids)
    knowledge_ids = {chunk.knowledge_id for chunk in chunks}
    knowledge_map = get_knowledge_by_ids(db, knowledge_ids)
    chunk_map = {c.vector_id: c for c in chunks}
    final = []
    for r in results:
        chunk = chunk_map.get(r["id"])
        knowledge = knowledge_map.get(chunk.knowledge_id) if chunk else None
        if chunk and knowledge and knowledge.is_enabled != 0:
            final.append({
                "chunk_id": chunk.id,
                "content": chunk.content,
                "score": max(0.0, 1.0 - float(r["distance"])),
                "distance": r["distance"],
                "knowledge_id": chunk.knowledge_id,
                "chunk_index": chunk.chunk_index,
            })
    if not final:
        logger.info(f"向量命中但未在MySQL找到chunk: agent={agent_id}, vector_ids={vector_ids[:5]}")
        return []

    reranker = _get_rerank_client()
    if reranker and len(final) > 1:
        doc_contents = [item["content"] for item in final]
        reranked = reranker.rerank(query, doc_contents, top_n=len(final))
        new_final = []
        for original_idx, rerank_score in reranked:
            if original_idx < len(final):
                item = final[original_idx].copy()
                item["score"] = rerank_score
                new_final.append(item)
        final = new_final[:top_k]
        logger.info(
            f"Rerank重排完成: {len(doc_contents)}条 → {len(final)}条, "
            f"Rerank最高分数={final[0]['score']:.4f}"
        )

    logger.info(f"检索完成: agent={agent_id}, query='{query[:20]}...', 命中{len(final)}条")
    for item in final:
        knowledge = knowledge_map.get(item["knowledge_id"])
        item["file_name"] = knowledge.file_name if knowledge else ""
        item["file_type"] = knowledge.file_type if knowledge else ""
    return final


async def search_async(
        db, user_id: int, agent_id: int, query: str,
        top_k: int = 5, knowledge_id: int = None,
) -> List[Dict[str, Any]]:
    """彻底 async 版检索：`db` 为 AsyncSession。

    向量化走 async 客户端 + async 配置查询；chunk / knowledge 反查走 async DAO；
    ChromaDB 相似度检索与 rerank 仍同步，统一封 `asyncio.to_thread`。
    行为与同步 `search` 逐条对齐。
    """
    if not query or not query.strip():
        raise ValueError("检索关键词不能为空")
    query_vector = await aembed_query_async(db, user_id, query)
    return await _build_search_results_async(db, agent_id, query, top_k, query_vector, knowledge_id)


async def _build_search_results_async(
        db, agent_id: int, query: str, top_k: int,
        query_vector: List[float], knowledge_id: int = None,
) -> List[Dict[str, Any]]:
    """`_build_search_results` 的 AsyncSession 版。"""
    import asyncio

    from models.knowledge_async_dao import (
        get_chunks_by_vector_ids_async, get_knowledge_by_ids_async,
    )

    if not query_vector:
        return []
    logger.info(f"开始检索(async): agent={agent_id}, query='{query[:50]}...', top_k={top_k}")
    retrieve_count = max(10, top_k * 2) if RAG_RERANK_ENABLED else top_k
    where = {"knowledge_id": knowledge_id} if knowledge_id is not None else None
    results = await asyncio.to_thread(
        search_similar, agent_id, query_vector, retrieve_count, where,
    )
    if not results:
        logger.info(f"向量检索无命中(async): agent={agent_id}, query='{query[:50]}...'")
        return []
    logger.info(f"向量检索完成(async): agent={agent_id}, 命中{len(results)}条")

    vector_ids = [r["id"] for r in results]
    chunks = await get_chunks_by_vector_ids_async(db, vector_ids)
    knowledge_ids = {chunk.knowledge_id for chunk in chunks}
    knowledge_map = await get_knowledge_by_ids_async(db, knowledge_ids)
    chunk_map = {c.vector_id: c for c in chunks}
    final = []
    for r in results:
        chunk = chunk_map.get(r["id"])
        knowledge = knowledge_map.get(chunk.knowledge_id) if chunk else None
        if chunk and knowledge and knowledge.is_enabled != 0:
            final.append({
                "chunk_id": chunk.id,
                "content": chunk.content,
                "score": max(0.0, 1.0 - float(r["distance"])),
                "distance": r["distance"],
                "knowledge_id": chunk.knowledge_id,
                "chunk_index": chunk.chunk_index,
            })
    if not final:
        logger.info(f"向量命中但未在MySQL找到chunk(async): agent={agent_id}, vector_ids={vector_ids[:5]}")
        return []

    reranker = _get_rerank_client()
    if reranker and len(final) > 1:
        doc_contents = [item["content"] for item in final]
        reranked = await asyncio.to_thread(reranker.rerank, query, doc_contents, len(final))
        new_final = []
        for original_idx, rerank_score in reranked:
            if original_idx < len(final):
                item = final[original_idx].copy()
                item["score"] = rerank_score
                new_final.append(item)
        final = new_final[:top_k]
        logger.info(
            f"Rerank重排完成(async): {len(doc_contents)}条 → {len(final)}条, "
            f"Rerank最高分数={final[0]['score']:.4f}"
        )

    logger.info(f"检索完成(async): agent={agent_id}, query='{query[:20]}...', 命中{len(final)}条")
    for item in final:
        knowledge = knowledge_map.get(item["knowledge_id"])
        item["file_name"] = knowledge.file_name if knowledge else ""
        item["file_type"] = knowledge.file_type if knowledge else ""
    return final
# ========== 彻底删除文档 ==========
def delete_knowledge_completely(
        db,agent_id:int,knowledge_id:int
)->Dict[str,Any]:
    """
    彻底删除文档,按三层顺序删。先删最危险的（容易漏的向量库），再往内一层层删。
    """
    knowledge = get_knowledge_by_id(db,knowledge_id)
    if not knowledge:
        return {"message": "文档不存在"}
    # 1. 删向量库（即使失败也继续，保证DB记录被清）
    try:
        delete_vectors_by_knowledge(_vector_key(knowledge),knowledge_id)
    except Exception as e:
        logger.warning(f"删向量库失败（继续删DB）: {e}")
    # 2. 删 chunks
    deleted_chunks=delete_chunks_by_knowledge(db,knowledge_id)
    # 3. 删 knowledge 主记录
    dao_delete_knowledge(db,knowledge)
    logger.info(
        f"文档彻底删除: knowledge_id={knowledge_id}, 删除chunks={deleted_chunks}"
    )
    return {"message": "删除成功", "deleted_chunks": deleted_chunks}
